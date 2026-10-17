import json,sys;sys.path.insert(0,'/home/claude/w')
from ld import *
import numpy as np
REP={'Abdullah Hussain I Khubrani':'Abdullah','DAI FAHAD A ALOTAIBI':'Dhai','Razan Areshyi':'Razan','Saleh Ali S Al Shehri':'Saleh','Vacant Azi Once - R1':'Vacant Azi-Once','Vacant Ciproxen - R1':'Vacant Ciproxen','VACCANT NON - Promoted R1':'Vacant Non-Promoted','R1-Non Promoted':'Vacant Non-Promoted','Vacant Prima - R1':'Vacant Prima'}
def mm(df,col,yr):
    return df[df.Year==yr].groupby('Period')[col].sum().reindex(range(1,13)).fillna(0).round(0).tolist()
def R(x): return float(round(x))
def build(o,tgt,fc):
    o=o.copy();o['cid']=o['New Custid'].astype(str);o['rep']=o.MED_REP_NAME.map(REP)
    tgt=tgt.copy();tgt['cid']=tgt['New Custid'].astype(str)
    sal=o[o.rep=='Saleh']
    out={}
    out['team']=[{'rep':r,'m26':mm(o[o.rep==r],'NET Value',2026),'m25':mm(o[o.rep==r],'NET Value',2025),'b26':mm(o[o.rep==r],'BUDGET',2026)} for r in dict.fromkeys(REP.values())]
    prods=[]
    for p in list(dict.fromkeys(sal.Product)):
        d=sal[sal.Product==p]
        prods.append({'p':p,'m26':mm(d,'NET Value',2026),'m25':mm(d,'NET Value',2025),'b26':mm(d,'BUDGET',2026)})
    out['prods']=prods
    return out,o,sal

DAL={'hosp':'105301310003','care':'435040','cp':'539051','nmr':'105006'}
def ytd(d):
    return [R(d[d.Year==2026]['NET Value'].sum()),R(d[(d.Year==2025)&(d.Period<=9)]['NET Value'].sum())]
def lst(a): return [float(x) for x in a]
def build_full(o,tgt,fc):
    out,o,sal=build(o,tgt,fc)
    tgt['cid']=tgt['New Custid'].astype(str)
    fcs=fc[fc['Medical rep name'].astype(str).str.contains('Saleh')].copy()
    fcs['cid']=fcs['Cust - ID'].astype(str)
    q=fcs[fcs.Period_ID.isin([10,11,12])].copy()
    q['lv']=q['Forecast Units']*q['X - Factor']
    # extra products from forecast with no sales rows
    names=[p['p'] for p in out['prods']]
    for p in dict.fromkeys(q.Product):
        if p not in names: out['prods'].append({'p':p,'m26':[0.0]*12,'m25':[0.0]*12,'b26':[0.0]*12})
    # prices
    price={}
    for p,d in q.groupby('Product',sort=False):
        dd=d[d['Forecast Units']>0]
        if len(dd) and dd['Forecast Units'].sum()>0:
            netw=((dd['NET - Price']*dd['Forecast Units']).sum()/dd['Forecast Units'].sum())
            price[p]={'list':round(float(dd['X - Factor'].max()),2),'net_w':round(float(netw),2),'net_min':round(float(dd['NET - Price'].min()),2),'net_max':round(float(dd['NET - Price'].max()),2)}
    out['price']=price
    cust=[]
    for i in dict.fromkeys(sal.cid):
        if i=='Budget': continue
        d=sal[sal.cid==i]
        m26=mm(d,'NET Value',2026);m25=mm(d,'NET Value',2025);tm=tgt[(tgt.cid==i)&(tgt.MED_REP_NAME.astype(str).str.contains('Saleh'))].groupby('Period')['BUDGET'].sum().reindex(range(1,13)).fillna(0).round(0).tolist()
        y26,y25=ytd(d); q425=R(d[(d.Year==2025)&(d.Period>=10)]['NET Value'].sum())
        if not any([y26,y25,q425,sum(tm)]) : continue
        pr=d.groupby('Product').apply(lambda x:[R(x[x.Year==2026]['NET Value'].sum()),R(x[(x.Year==2025)&(x.Period<=9)]['NET Value'].sum())])
        prods=sorted([[k,v[0],v[1]] for k,v in pr.items() if abs(v[0])+abs(v[1])>0],key=lambda z:-z[1])
        fq=q[q.cid==i]
        lv=fq['lv'].sum(); disc=round(1-fq['Net Value'].sum()/lv,3) if lv>0 else None
        cust.append({'id':i,'n':d['Customer Name'].iloc[0],'nat':d.NATURE.iloc[0],'sub':d['SUB - NATURE'].iloc[0],'y26':y26,'y25':y25,'q425':q425,'m26':m26,'m25':m25,'tm':tm,'fcm':[0.0]*3,'fc':R(fq['Net Value'].sum()),'disc':disc,'prods':prods,'fprods':[]})
    cust.sort(key=lambda c:-c['fc'])
    out['cust']=cust
    out['nat']=[[k,v[0],v[1]] for k,v in sal[sal['SUB - NATURE']!='Budget - 26'].groupby('SUB - NATURE').apply(ytd).items()]
    # lines
    lines=[]
    for ln in sorted(o.Line.dropna().unique()):
        d=o[o.Line==ln]
        grp=[]
        for g in d['PRODUCT - GROUP'].dropna().unique():
            dg=d[d['PRODUCT - GROUP']==g]
            ps=[]
            for p in dg.Product.unique():
                dp=dg[dg.Product==p]; y=ytd(dp)
                ps.append({'p':p,'y26':y[0],'y25':y[1],'b':R(dp[(dp.Year==2026)&(dp.Period<=9)]['BUDGET'].sum()),'fb':R(dp[dp.Year==2026]['BUDGET'].sum()),'q425':R(dp[(dp.Year==2025)&(dp.Period>=10)]['NET Value'].sum())})
            y=ytd(dg)
            grp.append({'g':g,'y26':y[0],'y25':y[1],'b':R(dg[(dg.Year==2026)&(dg.Period<=9)]['BUDGET'].sum()),'fb':R(dg[dg.Year==2026]['BUDGET'].sum()),'q425':R(dg[(dg.Year==2025)&(dg.Period>=10)]['NET Value'].sum()),'bq4':R(dg[(dg.Year==2026)&(dg.Period>=10)]['BUDGET'].sum()),'prods':ps})
        dd=d[d['SUB - NATURE']!='Budget - 26']
        lines.append({'line':ln,'reps':list(dict.fromkeys(d.rep)),'m26':mm(d,'NET Value',2026),'m25':mm(d,'NET Value',2025),'b26':mm(d,'BUDGET',2026),'groups':grp,'nat':[[k,v[0],v[1]] for k,v in dd.groupby('SUB - NATURE').apply(ytd).items()],'strat':[[k,v[0]] for k,v in d[d['2026 Strategic Type'].notna()].groupby('2026 Strategic Type').apply(ytd).items()]})
    out['lines']=lines
    dal={}
    for k,c in DAL.items():
        d=o[o.cid==c]
        byrep={r:ytd(d[d.rep==r]) for r in dict.fromkeys(REP.values()) if (d.rep==r).any()}
        pp=d.groupby('Product').apply(ytd)
        byprod=sorted([[p,v[0],v[1]] for p,v in pp.items() if abs(v[0])+abs(v[1])>0],key=lambda z:-z[1])
        dal[k]={'m26':mm(d,'NET Value',2026),'m25':mm(d,'NET Value',2025),'byrep':byrep,'byprod':byprod}
    dal['share26']={r:R(sum(dal[k]['byrep'].get(r,[0,0])[0] for k in DAL)) for r in dict.fromkeys(REP.values())}
    out['dallah']=dal
    return out

def mkcube(o,t,H_old):
    o=o.copy();t=t.copy()
    for d in(o,t):
        d['cid']=d['New Custid'].astype(str);d['rep']=d.MED_REP_NAME.map(REP);d['pid']=d.Productid.astype(str)
    o=o[o.cid!='Budget'];t=t[t.cid!='Budget']
    k=['cid','pid','rep']
    def agg(df,mask,col,name): return df[mask].groupby(k)[col].sum().rename(name)
    parts=[agg(o,(o.Year==2026)&(o.Period<=9),'NET Value','y26'),agg(o,(o.Year==2025)&(o.Period<=9),'NET Value','y25'),agg(o,(o.Year==2025)&(o.Period>=10),'NET Value','q4'),agg(o,(o.Year==2025),'NET Value','fy25'),
           agg(t,t.Period<=9,'BUDGET','bY'),agg(t,t.Period>=1,'BUDGET','bF'),agg(t,t.Period>=10,'BUDGET','bQ')]
    df=pd.concat(parts,axis=1).fillna(0).round(0)
    df=df[df.abs().sum(axis=1)>0].reset_index()
    # dims
    reps=list(dict.fromkeys(REP.values()))
    pm=o.drop_duplicates('pid').set_index('pid')
    pm2=t.drop_duplicates('pid').set_index('pid')
    pids=list(dict.fromkeys(df.pid))
    # keep old ordering for known products, append new
    oldp=[p[0] for p in H_old['dims']['p']]
    pids=[p for p in oldp if p in set(pids)]+[p for p in pids if p not in set(oldp)]
    def prow(p):
        r=pm.loc[p] if p in pm.index else pm2.loc[p]
        return [p,str(r['Product']),str(r['PRODUCT - GROUP']),str(r['Line']),str(r['Category']),str(r['2026 Strategic Type'])]
    dp=[prow(p) for p in pids]
    cm=o.drop_duplicates('cid').set_index('cid');cm2=t.drop_duplicates('cid').set_index('cid')
    cids=list(dict.fromkeys(df.cid));oldc=[c[0] for c in H_old['dims']['c']]
    cids=[c for c in oldc if c in set(cids)]+[c for c in cids if c not in set(oldc)]
    def crow(c):
        r=cm.loc[c] if c in cm.index else cm2.loc[c]
        return [c,str(r['Customer Name']),str(r['SUB - NATURE']),str(r['NATURE'])]
    dc=[crow(c) for c in cids]
    pi={p:i for i,p in enumerate(pids)};ci={c:i for i,c in enumerate(cids)};ri={r:i for i,r in enumerate(reps)}
    rows=[[ci[r.cid],pi[r.pid],ri[r.rep],r.y26,r.y25,r.q4,r.fy25,r.bY,r.bF,r.bQ] for r in df.itertuples()]
    return {'p':dp,'c':dc,'r':reps},[[int(x[0]),int(x[1]),int(x[2])]+[float(v) for v in x[3:]] for x in rows]

def mk_habib(o,dims,rows,old_h):
    import openpyxl
    y=openpyxl.load_workbook('/mnt/user-data/uploads/Habib_Sales_Out___YTG_-_Plan__4_.xlsx',data_only=True).worksheets[0]
    ytg=[r for r in y.iter_rows(min_row=2,values_only=True) if r[0]]
    wb=openpyxl.load_workbook(glob.glob(NEW+'*3ce62b13*')[0],data_only=True)['Sales Data']
    drows=list(wb.iter_rows(min_row=3,values_only=True))
    cat=None;D={}
    for r in drows:
        if not r[0]: continue
        if r[0] in('Promotion','Commercial'): cat=r[0];continue
        if r[0]=='Grand Total': continue
        D[r[0]]={'cat':cat,'um26':[float(x or 0) for x in r[1:10]],'um25':[float(x or 0) for x in r[11:20]]}
    o=o.copy();o['cid']=o['New Custid'].astype(str);o['pid']=o.Productid.astype(str)
    h=o[o.cid=='105301120009']
    pmap={p[0]:p for p in dims['p']};reps=dims['r']
    oldp={p['code']:p for p in old_h['prods']}
    prods=[]
    for r in ytg:
        code=str(r[0]);name=r[2];price=float(r[4])
        d=h[h.pid==code]
        si26=[R(x) for x in d[d.Year==2026].groupby('Period')['NET Value'].sum().reindex(range(1,10)).fillna(0)]
        si25=[R(x) for x in d[d.Year==2025].groupby('Period')['NET Value'].sum().reindex(range(1,13)).fillna(0)]
        x=D[name];u=sum(x['um26']);so=R(u*price);si=R(sum(si26))
        f=[float(v or 0) for v in r[9:12]]  # Oct,Nov,Dec
        pr=pmap.get(code,[code,name,'','','',''])
        ri=[]
        for rr in rows:
            if dims['p'][rr[1]][0]==code and (rr[3]!=0 or rr[7]>0): 
                rn=reps[rr[2]]
                if rn not in ri: ri.append(rn)
        op=oldp.get(code,{})
        prods.append({'code':code,'name':name,'line':pr[3],'brand':pr[2],'st':pr[5],'cat':x['cat'],'reps':ri,'price':price,'u':u,'so':so,'si':si,'dd':so-si,
           'f':f,'fu':round(sum(f)),'fv':R(sum(f)*price),'um26':x['um26'],'um25':x['um25'],'si26':si26,'si25':si25,'d26':op.get('d26',[]),'d25':op.get('d25',[]),'hc':str(r[1])})
    wm=h
    m={'si26':[R(v) for v in wm[wm.Year==2026].groupby('Period')['NET Value'].sum().reindex(range(1,10)).fillna(0)],'si25':[R(v) for v in wm[wm.Year==2025].groupby('Period')['NET Value'].sum().reindex(range(1,13)).fillna(0)],'obj':1591000}
    return {'prods':prods,'m':m,'id':'105301120009'}

def mk_fcx(fc,old_rows,price):
    f=fc[fc['Medical rep name'].astype(str).str.contains('Saleh')&fc.Period_ID.isin([10,11,12])].copy()
    f['cid']=f['Cust - ID'].astype(str)
    f=f[f.cid!='Budget - 26']
    tot_by=f.groupby('cid')['Forecast Units'].sum();oldc={r[0] for r in old_rows}
    f=f[f.cid.map(lambda c: tot_by[c]>0 or c in oldc)]
    rows=[];seen=set()
    for (c,p),d in f.groupby(['cid','Product'],sort=False):
        m=[float(d[d.Period_ID==k]['Forecast Units'].sum()) for k in (10,11,12)]
        tot=sum(m)
        if p=='EMEDIUS INJECTION' and tot==0: continue
        if tot>0:
            pr=float(d['Net Value'].sum()/tot)
        else:
            pr=float(price.get(p,{}).get('net_w') or d['NET - Price'].dropna().mean() or 0)
        rows.append([c,p,round(pr,4),m[0],m[1],m[2]]);seen.add((c,p))
    for r in old_rows:
        if (r[0],r[1]) not in seen: rows.append([r[0],r[1],r[2],0.0,0.0,0.0])
    return rows
