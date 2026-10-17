import sys,json,re,datetime;sys.path.insert(0,'/home/claude/w')
from mk_sales import *
import openpyxl
P=json.load(open('/home/claude/w/parts_old.json'))
nsales=sales('new');nt=ld(glob.glob(NEW+'*c7594402*')[0],'26 - Target Per Account')
nf=ld(glob.glob(NEW+'*6234ae5b*')[0],'Q4 - Months')
N=build_full(nsales,nt,nf)
# keep forecast-derived keys from old (runtime recomputes) 
for k in ('fdal','fprod'):
    N[k]=P['SALES'][k]
dims,rows=mkcube(nsales,nt,P['HAB'])
hb=mk_habib(nsales,dims,rows,P['HAB']['habib'])
HAB={'dims':dims,'cube':rows,'habib':hb}
fcx=mk_fcx(nf,P['FCX']['rows'],N['price'])
# runtime fills prods fc/fu: give placeholders to keep shape
for p in N['prods']:
    p['fc']=[0.0,0.0,0.0];p['fu']=[0.0,0.0,0.0]
# ---- branches
ws=openpyxl.load_workbook(glob.glob(NEW+'*a734fea4*')[0],data_only=True).worksheets[0]
rr=list(ws.iter_rows(values_only=True))
BN=[('ghadeer','Al-Ghadeer Polyclinic','الغدير','Suliman alhabib alghadeer'),('hammra','Al-Hammra Hospital','الحمراء','SOLAIMAN HABIB HOSPITAL Al Hamra'),('sweedie','Al-Sweedie Hospital','السويدي','SOLAIMAN HABIB HOSPITAL AL SWEEDY')]
hnames={p['name']:p['code'] for p in hb['prods']}
def nrm(s): return re.sub(r"[^a-z0-9]","",s.lower().replace('(dis)',''))
hn={nrm(k):v for k,v in hnames.items()}
bprods=[]
for r in rr[2:]:
    if not r[0] or r[0]=='Grand Total': continue
    vals=[[float(x or 0) for x in r[1+18*b:19+18*b]] for b in range(3)]
    if not any(any(v) for v in vals): continue
    bprods.append({'n':r[0],'code':hn.get(nrm(r[0]),''),'v':vals})
gt=[r for r in rr if r[0]=='Grand Total'][0]
tot=[[float(x or 0) for x in gt[1+18*b:19+18*b]] for b in range(3)]
chk=[[round(sum(p['v'][b][i] for p in bprods)-tot[b][i],1) for i in range(18)] for b in range(3)]
print('branch grand-total mismatches',[x for c in chk for x in c if abs(x)>1])
BR={'months':['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep'],'branches':[{'id':b[0],'n':b[1],'ar':b[2],'hf':b[3]} for b in BN],'prods':bprods,'tot':tot}
# ---- Jamjoom stock
j=pd.read_excel(glob.glob(NEW+'*6a84f533*')[0])
brandmap={'ACILOC':'Aciloc','AZI-ONCE':'Azi-Once','BETAVERT':'Betavert','CARTIMOV':'Cartimov','CIPROXEN':'Ciproxen','CLARA':'Clara','DESLOR':'Deslor','DOMPY':'Dompy','ERECTA':'Erecta','FENSOLIN':'Fensolin','FEROTON':'Feroton','HEPAFORM':'Hepaform','MAXIM':'Maxim','MEVA':'Meva','OSELTA':'Oselta','PRIMA':'PRIMA','PILOCIN':'Pilocin','RELAXON':'Relaxon','TORICOX':'Toricox','TRIPLEX':'Triplex-B','URAGO':'Urago','VOLTIC':'Voltic','YOCAIR':'Yocair','ZORON':'Zoron'}
b2l={}
for p in dims['p']: b2l.setdefault(p[2],p[3])
JJrows=[]
for r in j.itertuples(index=False):
    d=str(r[1]).strip();d2=re.sub(r'\s+',' ',d)
    words=re.sub(r'^(W-|JP )','',d2).split(' ')
    key=words[0].upper().strip(',');br=brandmap.get(key,'')
    if key=='JP' or d2.startswith('JP TRIPLEX') or d2.startswith('W-JP TRIPLEX'): br='Triplex-B'
    JJrows.append({'c':int(r[0]),'d':d2,'p':float(r[2]),'wh':int(r[3]),'ph':int(r[4]),'tr':int(r[5]),'rt':int(r[6]),'tot':int(r[7]),'s30':int(r[8]),'s60':int(r[9]),'m':str(r[10]),'b':br,'l':b2l.get(br,'')})
JJ={'date':'2026-10-04','rows':JJrows,'map':{"MAXIM 400MG TAB 1X5":67238,"MAXIM 400MG TAB 1x7":67239,"MEVA 135 MG F.C. TABS X 30":18445,"MEVA 200 MG F.C. TABS X 30":95939,"ZORON 4Mg Tablet 10's":93597,"ZORON 8Mg Tablet 10's":93598,"ZORON Injection 8Mg/2ml 5's":88463}}
# ---- previous version snapshot for diff
PREV={'label':'Version of 13 Oct (Q4 forecast R1 · sales through Sep, earlier cut)','fcx':P['FCX']['rows'],
 'team':{t['rep']:round(sum(t['m26'][:9])) for t in P['SALES']['team']},
 'sep':{t['rep']:round(t['m26'][8]) for t in P['SALES']['team']},
 'o':0.9584,'h':59.7,'fcsar':2594570,'ytd':2513436,
 'habib':{'u':round(sum(sum(p['um26']) for p in P['HAB']['habib']['prods'])),'so':round(sum(p['so'] for p in P['HAB']['habib']['prods'])),'si':round(sum(p['si'] for p in P['HAB']['habib']['prods']))}}
out={'D2':P['D2'],'SALES':N,'HAB':HAB,'HDOC':P['HDOC'],'TERR':P['TERR'],'FCX':{'rows':fcx},'BR':BR,'JJ':JJ,'PREV':PREV}
json.dump(out,open('/home/claude/w/parts_new.json','w'))
print({k:len(json.dumps(v)) for k,v in out.items()})
print(sum(1 for r in JJrows if r['b']==''),[r['d'] for r in JJrows if r['b']==''][:60])
print('Saleh YTD',sum(sum(p['m26'][:9]) for p in N['prods']),'fcx',sum(r[3]+r[4]+r[5] for r in fcx))
print('branch YTD26 units', [round(sum(t[9:18])) for t in tot],'habib total',round(sum(sum(p['um26']) for p in hb['prods'])))
